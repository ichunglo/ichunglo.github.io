"""Import the approved-drug supplements into normalized CSV/JSON tables.

This is a reproducible Phase 1 import, not a clinical-program curation script.
The original XLSX files remain external source material. The generated tables
retain paper/source names and deterministic IDs for later reconciliation.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
from datetime import date
from pathlib import Path

from openpyxl import load_workbook


SOURCE_DIR = Path(r"C:\Users\ichun\Documents\Codex\The Envolving Landscape of Drug Targets")
ROOT = Path(__file__).resolve().parents[1]
CURATED = ROOT / "data" / "curated"
GENERATED = ROOT / "data" / "generated"
MANIFEST = ROOT / "data" / "raw-manifest.csv"

AS_OF_DATE = "2026-10-02"
RELEASE_VERSION = "0.1.0"
CGT_ONLY_GENES = {"CNTF", "CSF2", "IFNA2", "ACPP", "ARSA", "RPE65", "F8", "BCL11A", "COL7A1", "HBB", "MAGEA4", "WAS", "ABCD1", "OTOF"}


def slug(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9]+", "-", value.strip().lower()).strip("-")
    return value or "unknown"


def normalized_name(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.strip().lower())


def split_cell(value: object, delimiter: str = ",") -> list[str]:
    if value is None:
        return []
    text = str(value).strip()
    if not text:
        return []
    return [part.strip() for part in text.split(delimiter) if part.strip()]


def read_sheet(filename: str, sheet: str) -> list[dict[str, object]]:
    path = SOURCE_DIR / filename
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet]
    rows = worksheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    return [dict(zip(headers, row)) for row in rows]


def write_csv(path: Path, rows: list[dict[str, object]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def main() -> None:
    human = read_sheet("41573_2026_1530_MOESM2_ESM.xlsx", "Human_Drug_Targets")
    pathogen = read_sheet("41573_2026_1530_MOESM3_ESM.xlsx", "Non_Human_Drug_Targets")
    approvals = read_sheet("41573_2026_1530_MOESM4_ESM.xlsx", "Drug-Approval-Target-ATC")
    kinases = read_sheet("41573_2026_1530_MOESM5_ESM.xlsx", "Kinase_Drug_Targets")

    organisms: dict[str, dict[str, object]] = {}
    targets: dict[str, dict[str, object]] = {}
    drugs: dict[str, dict[str, object]] = {}
    drug_targets: dict[tuple[str, str], dict[str, object]] = {}
    aliases: dict[tuple[str, str], dict[str, object]] = {}
    approval_events: list[dict[str, object]] = []

    def add_organism(name: str, scope: str) -> int:
        key = name.lower()
        if key not in organisms:
            organisms[key] = {"organism_id": len(organisms) + 1, "name": name, "taxonomy_scope": scope}
        return int(organisms[key]["organism_id"])

    def add_target(name: str, gene: str | None, accession: str | None, organism_name: str, scope: str, group: str | None, source_row: str, landscape_scope: str | None = None) -> str:
        if name == "Tubulin beta chain":
            name = "Tubulin beta"
        genes = set(re.split(r"[|;]", gene or ""))
        target_scope = landscape_scope or ("cgt_only" if genes & CGT_ONLY_GENES else "core")
        organism_id = add_organism(organism_name, scope)
        accession_key = (accession or "").lower()
        if name == "Tubulin beta":
            key = f"tubulin-beta|{organism_id}"
        else:
            key = f"{accession_key}|{organism_id}" if accession_key else f"{organism_id}|{slug(name)}|{source_row if scope == 'pathogen' else ''}"
        target_suffix = slug(source_row) if not accession_key else ""
        target_id = f"TGT-{slug(accession or name)}-{organism_id}-{target_suffix}".rstrip("-")
        if key not in targets:
            targets[key] = {
                "target_id": target_id,
                "preferred_name": name,
                "gene_symbol": gene or "",
                "uniprot_accession": accession or "",
                "organism_id": organism_id,
                "target_status": "confirmed" if accession else "generalized_pathogen" if scope == "pathogen" else "unclear",
                "landscape_scope": target_scope,
                "protein_group": group or "",
                "target_type": "",
                "complex_description": "",
                "paper_target_name": name,
                "paper_source_row": source_row,
            }
        return str(targets[key]["target_id"])

    def find_or_add_name_target(name: str) -> str:
        needle = normalized_name(name)
        for record in targets.values():
            if normalized_name(str(record["preferred_name"])) == needle:
                return str(record["target_id"])
        return add_target(name, None, None, "human", "other", None, "MOESM4", "unresolved")

    def add_drug(name: str) -> str:
        drug_id = f"DRG-{slug(name)}"
        if drug_id not in drugs:
            drugs[drug_id] = {
                "drug_id": drug_id,
                "generic_name": name,
                "preferred_name": name,
                "modality": "unknown",
                "chemical_or_biologic_id": "",
                "cas_registration_year": "",
                "atc_code": "",
                "source_name": "The evolving landscape of drug targets",
            }
        return drug_id

    def link(drug_id: str, target_id: str, role: str, source_id: str) -> None:
        key = (drug_id, target_id)
        if key not in drug_targets:
            drug_targets[key] = {
                "drug_target_id": f"DT-{slug(drug_id)}-{slug(target_id)}",
                "drug_id": drug_id,
                "target_id": target_id,
                "target_role": role,
                "moa_status": "confirmed" if role == "primary" else "secondary",
                "potency_value": "",
                "potency_unit": "",
                "potency_log10_molar": "",
                "evidence_grade": "paper-derived",
                "source_id": source_id,
            }

    for index, row in enumerate(human, start=2):
        if not row["TARGET_PREF_NAME"]:
            continue
        target_id = add_target(str(row["TARGET_PREF_NAME"]), str(row["GENE_NAME"] or ""), str(row["ACCESSION"] or ""), "human", "human", str(row["PROTEIN_GROUP"] or ""), f"MOESM2:{index}")
        for drug_name in split_cell(row["DRUG_NAMES (1-3)"]):
            link(add_drug(drug_name), target_id, "primary", "SRC-MOESM2")

    for index, row in enumerate(pathogen, start=2):
        if not row["TARGET_PREF_NAME"]:
            continue
        organism = str(row["ORGANISM"] or "other")
        target_id = add_target(str(row["TARGET_PREF_NAME"]), None, None, organism, "pathogen", None, f"MOESM3:{index}")
        for drug_name in split_cell(row["DRUG_NAMES (1-3)"]):
            link(add_drug(drug_name), target_id, "primary", "SRC-MOESM3")

    for index, row in enumerate(kinases, start=2):
        if not row["TARGET_PREF_NAME"]:
            continue
        target_id = add_target(str(row["TARGET_PREF_NAME"]), str(row["GENE_NAME"] or ""), str(row["ACCESSION"] or ""), "human" if str(row["GENE_NAME"] or "") else "other", "human", str(row["PROTEIN_GROUP"] or ""), f"MOESM5:{index}")
        for drug_name in split_cell(row["DRUG_NAMES (1-3)"]):
            link(add_drug(drug_name), target_id, "primary", "SRC-MOESM5")

    for index, row in enumerate(approvals, start=2):
        drug_name = str(row["DRUG_NAME"] or "").strip()
        if not drug_name:
            continue
        drug_id = add_drug(drug_name)
        year = row["First Approval Year"]
        year_text = str(year).strip() if year is not None else ""
        approval_date = f"{int(year)}-01-01" if isinstance(year, (int, float)) and not isinstance(year, bool) else ""
        atc = str(row["ATC_CODE"] or "").strip()
        drugs[drug_id]["cas_registration_year"] = drugs[drug_id]["cas_registration_year"] or ""
        drugs[drug_id]["atc_code"] = atc or drugs[drug_id]["atc_code"]
        targets_in_row = split_cell(row["Target name"], ";")
        for position, target_name in enumerate(targets_in_row):
            target_id = find_or_add_name_target(target_name)
            link(drug_id, target_id, "primary" if position == 0 else "secondary", "SRC-MOESM4")
            approval_events.append({
                "approval_event_id": f"APR-{slug(drug_name)}-{slug(target_name)}-{slug(year_text or 'unknown')}-{index}",
                "drug_id": drug_id,
                "agency": "paper-derived",
                "country_or_region": "global/unspecified",
                "indication_id": "",
                "approval_date": approval_date,
                "approval_type": "first approval year",
                "source_id": "SRC-MOESM4",
            })

    source_rows = [
        {"source_id": "SRC-MOESM2", "source_type": "supplementary_workbook", "title": "Human drug targets", "publisher": "Nature Reviews Drug Discovery", "url": "https://doi.org/10.1038/s41573-026-01530-3", "retrieval_date": AS_OF_DATE, "citation": "41573_2026_1530_MOESM2_ESM.xlsx", "source_priority": 1, "notes": "Paper-derived baseline"},
        {"source_id": "SRC-MOESM3", "source_type": "supplementary_workbook", "title": "Non-human drug targets", "publisher": "Nature Reviews Drug Discovery", "url": "https://doi.org/10.1038/s41573-026-01530-3", "retrieval_date": AS_OF_DATE, "citation": "41573_2026_1530_MOESM3_ESM.xlsx", "source_priority": 1, "notes": "Paper-derived baseline"},
        {"source_id": "SRC-MOESM4", "source_type": "supplementary_workbook", "title": "Drug approval, target, and ATC table", "publisher": "Nature Reviews Drug Discovery", "url": "https://doi.org/10.1038/s41573-026-01530-3", "retrieval_date": AS_OF_DATE, "citation": "41573_2026_1530_MOESM4_ESM.xlsx", "source_priority": 1, "notes": "Paper-derived baseline"},
        {"source_id": "SRC-MOESM5", "source_type": "supplementary_workbook", "title": "Kinase drug targets", "publisher": "Nature Reviews Drug Discovery", "url": "https://doi.org/10.1038/s41573-026-01530-3", "retrieval_date": AS_OF_DATE, "citation": "41573_2026_1530_MOESM5_ESM.xlsx", "source_priority": 1, "notes": "Paper-derived baseline"},
    ]

    fields = {
        "organisms": ["organism_id", "name", "taxonomy_scope"],
        "targets": ["target_id", "preferred_name", "gene_symbol", "uniprot_accession", "organism_id", "target_status", "landscape_scope", "protein_group", "target_type", "complex_description", "paper_target_name", "paper_source_row"],
        "drugs": ["drug_id", "generic_name", "preferred_name", "modality", "chemical_or_biologic_id", "cas_registration_year", "atc_code", "source_name"],
        "drug-targets": ["drug_target_id", "drug_id", "target_id", "target_role", "moa_status", "potency_value", "potency_unit", "potency_log10_molar", "evidence_grade", "source_id"],
        "approvals": ["approval_event_id", "drug_id", "agency", "country_or_region", "indication_id", "approval_date", "approval_type", "source_id"],
        "sources": ["source_id", "source_type", "title", "publisher", "url", "retrieval_date", "citation", "source_priority", "notes"],
    }
    rows = {
        "organisms": list(organisms.values()),
        "targets": list(targets.values()),
        "drugs": list(drugs.values()),
        "drug-targets": list(drug_targets.values()),
        "approvals": approval_events,
        "sources": source_rows,
    }
    for name, values in rows.items():
        write_csv(CURATED / f"{name}.csv", values, fields[name])

    source_files = [
        "The evolving landscape of Drug Targets.pdf",
        "41573_2026_1530_MOESM1_ESM.pdf",
        "41573_2026_1530_MOESM2_ESM.xlsx",
        "41573_2026_1530_MOESM3_ESM.xlsx",
        "41573_2026_1530_MOESM4_ESM.xlsx",
        "41573_2026_1530_MOESM5_ESM.xlsx",
    ]
    manifest = []
    for filename in source_files:
        path = SOURCE_DIR / filename
        manifest.append({"filename": filename, "path": str(path), "sha256": sha256(path), "retrieval_date": AS_OF_DATE})
    write_csv(MANIFEST, manifest, ["filename", "path", "sha256", "retrieval_date"])

    approval_drugs = len({str(row["DRUG_NAME"]).strip().lower() for row in approvals if row["DRUG_NAME"]})
    summary = {
        "release_version": RELEASE_VERSION,
        "as_of_date": AS_OF_DATE,
        "source": "The evolving landscape of drug targets, supplementary tables",
        "counts": {"organisms": len(organisms), "targets": len(targets), "core_targets": sum(v["landscape_scope"] == "core" for v in targets.values()), "cgt_only_targets": sum(v["landscape_scope"] == "cgt_only" for v in targets.values()), "unresolved_target_labels": sum(v["landscape_scope"] == "unresolved" for v in targets.values()), "drugs_with_any_source_link": len(drugs), "approval_table_rows": len(approvals), "approval_table_drugs": approval_drugs, "drug_target_relationships": len(drug_targets), "approval_events": len(approval_events)},
        "paper_reference_counts": {"reported_biomolecules": 686, "reported_drugs": 1702},
        "headline_count_rule": "primary drug-target relationships only",
        "clinical_programs": 0,
        "notes": ["Phase 1 approved baseline import; clinical pilot not yet populated.", "Names introduced by MOESM4 without an existing UniProt mapping are retained as unresolved labels rather than core targets.", "CGT-only targets remain available as a separate scope and are excluded from core paper-comparable counts."],
    }
    GENERATED.mkdir(parents=True, exist_ok=True)
    browser_payload = {
        "release": {"version": RELEASE_VERSION, "as_of_date": AS_OF_DATE, "source": summary["source"]},
        "metrics": summary["counts"],
        "targets": list(targets.values()),
        "drugs": list(drugs.values()),
        "drug_targets": list(drug_targets.values()),
        "approvals": approval_events,
        "sources": source_rows,
        "clinical_programs": [],
    }
    (GENERATED / "landscape.json").write_text(json.dumps(browser_payload, indent=2), encoding="utf-8")
    (GENERATED / "baseline-summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
